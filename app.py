
import os, sqlite3, tempfile, re, math
from flask import Flask, request, jsonify, render_template, session, redirect, url_for, flash
from openpyxl import load_workbook

APP_DIR = os.path.dirname(os.path.abspath(__file__))
DB = os.path.join(APP_DIR, "rate_master.db")
DEFAULT_XLSX = os.path.join(APP_DIR, "IPHW_RATE_MASTER.xlsx")
ADMIN_PASSWORD = os.environ.get("ADMIN_PASSWORD", "admin123")

app = Flask(__name__)
app.secret_key = os.environ.get("SECRET_KEY", "change-this-secret-key")

def db():
    c = sqlite3.connect(DB)
    c.row_factory = sqlite3.Row
    return c

def init_db():
    c=db()
    c.executescript("""
    CREATE TABLE IF NOT EXISTS settings(
      key TEXT PRIMARY KEY, value TEXT NOT NULL
    );
    CREATE TABLE IF NOT EXISTS zones(country TEXT PRIMARY KEY, zone TEXT);
    CREATE TABLE IF NOT EXISTS rates(
      calculator TEXT, weight_from REAL, weight_to REAL, zone TEXT, rate REAL,
      PRIMARY KEY(calculator, weight_from, weight_to, zone)
    );
    CREATE TABLE IF NOT EXISTS oda(
      id INTEGER PRIMARY KEY AUTOINCREMENT,
      country TEXT, city TEXT, postal_from INTEGER, postal_to INTEGER,
      oda_type TEXT
    );
    CREATE TABLE IF NOT EXISTS import_log(
      id INTEGER PRIMARY KEY AUTOINCREMENT, imported_at TEXT DEFAULT CURRENT_TIMESTAMP,
      filename TEXT, status TEXT, message TEXT
    );
    """)
    c.commit()
    c.close()

def num(v, default=0):
    try: return float(v)
    except: return default

def import_xlsx(path, filename="IPHW_RATE_MASTER.xlsx"):
    wb=load_workbook(path, data_only=True)
    c=db()
    try:
        # settings from Calculator A
        ws=wb["Calculatore -A"]
        settings={
            "exchange_rate": num(ws["B7"].value,123.65),
            "vat_rate": num(ws["B8"].value,0.15),
            "oda_fuel": num(ws["B9"].value,0.5325),
            "oda_base_usd": num(ws["B15"].value,25),
            "oda_perkg_usd": num(ws["B16"].value,0.5),
        }
        for k,v in settings.items():
            c.execute("INSERT INTO settings(key,value) VALUES(?,?) ON CONFLICT(key) DO UPDATE SET value=excluded.value",(k,str(v)))

        c.execute("DELETE FROM zones")
        ws=wb["Sheet2"]
        for r in range(5, ws.max_row+1):
            country=ws.cell(r,1).value
            zone=ws.cell(r,2).value
            if country and zone:
                c.execute("INSERT OR REPLACE INTO zones(country,zone) VALUES(?,?)",(str(country).strip(),str(zone).strip()))

        c.execute("DELETE FROM rates")
        for sheet,calc in [("Rate ; A","A"),("Rate; B","B")]:
            ws=wb[sheet]
            zones=[ws.cell(1,col).value for col in range(2,ws.max_column+1)]
            for r in range(2,9):
                label=ws.cell(r,1).value
                if not label: continue
                if str(label).startswith("101"): lo,hi=101,999999
                else:
                    m=re.match(r"(\d+)-(\d+)",str(label))
                    if not m: continue
                    lo,hi=int(m.group(1)),int(m.group(2))
                for i,z in enumerate(zones, start=2):
                    val=ws.cell(r,i).value
                    if z and isinstance(val,(int,float)):
                        c.execute("INSERT INTO rates VALUES(?,?,?,?,?)",(calc,lo,hi,str(z),float(val)))

        c.execute("DELETE FROM oda")
        ws=wb["Sheet1"]
        for r in range(9,ws.max_row+1):
            country=ws.cell(r,1).value
            city=ws.cell(r,2).value
            pf=ws.cell(r,3).value
            pt=ws.cell(r,4).value
            if not country: continue
            country=str(country).strip()
            city=str(city).strip() if city is not None else ""
            try: pf=int(float(pf)) if pf is not None else None
            except: pf=None
            try: pt=int(float(pt)) if pt is not None else None
            except: pt=None
            typ="City" if city and pf is None else "Postal" if pf is not None else "City"
            c.execute("INSERT INTO oda(country,city,postal_from,postal_to,oda_type) VALUES(?,?,?,?,?)",
                      (country,city,pf,pt,typ))
        c.execute("INSERT INTO import_log(filename,status,message) VALUES(?,?,?)",
                  (filename,"OK","Rate, zone and ODA master imported"))
        c.commit()
    except Exception as e:
        c.rollback()
        c.execute("INSERT INTO import_log(filename,status,message) VALUES(?,?,?)",(filename,"ERROR",str(e)))
        c.commit()
        raise
    finally:
        c.close()

def get_settings():
    c=db()
    d={r["key"]:float(r["value"]) for r in c.execute("SELECT key,value FROM settings")}
    c.close()
    return d

def calc(payload):
    weight=float(payload.get("weight",0))
    country=(payload.get("country") or "").strip()
    postal=(payload.get("postal") or "").strip()
    city=(payload.get("city") or "").strip()
    calc_type=(payload.get("calculator") or "A").upper()
    s=get_settings()
    if weight<=0: raise ValueError("Enter a valid weight.")
    c=db()
    zrow=c.execute("SELECT zone FROM zones WHERE lower(country)=lower(?)",(country,)).fetchone()
    if not zrow: c.close(); raise ValueError("Country not found.")
    zone=zrow["zone"]
    r=c.execute("""SELECT rate FROM rates WHERE calculator=? AND weight_from<=? AND weight_to>=? AND zone=?""",
                (calc_type,math.ceil(weight),math.ceil(weight),zone)).fetchone()
    if not r:
        c.close(); raise ValueError(f"No rate found for Calculator {calc_type}, zone {zone}, weight {math.ceil(weight)}.")
    perkg=r["rate"]
    rows=c.execute("SELECT * FROM oda WHERE lower(country)=lower(?)",(country,)).fetchall()
    oda=False; matched=[]
    try: p=int(postal) if postal else None
    except: p=None
    for row in rows:
        hit=False
        if row["oda_type"]=="City" and city and row["city"].lower()==city.lower(): hit=True
        if row["oda_type"]=="Postal" and p is not None and row["postal_from"]<=p<=row["postal_to"]: hit=True
        if hit: oda=True; matched.append(dict(row))
    c.close()
    base=math.ceil(weight)*perkg
    oda_usd=max(s.get("oda_base_usd",25), math.ceil(weight)*s.get("oda_perkg_usd",.5)) if oda else 0
    oda_bdt=oda_usd*s.get("exchange_rate",0)
    oda_fuel=oda_bdt*s.get("oda_fuel",0) if oda else 0
    subtotal=base+oda_bdt+oda_fuel
    total=subtotal*(1+s.get("vat_rate",0))
    return {"calculator":calc_type,"weight":weight,"rounded_weight":math.ceil(weight),"country":country,
            "zone":zone,"per_kg_rate":perkg,"base_tariff":base,"oda_status":"ODA" if oda else "No ODA",
            "oda_usd":oda_usd,"oda_bdt":oda_bdt,"oda_fuel_bdt":oda_fuel,
            "subtotal_without_vat":subtotal,"vat_rate":s.get("vat_rate",0),
            "grand_total_with_vat":total,"exchange_rate":s.get("exchange_rate",0),
            "matched_by":sorted(set(x["oda_type"] for x in matched)),"matched_rows":matched[:50]}

@app.route("/")
def index():
    return render_template("index.html")

@app.route("/api/meta")
def meta():
    c=db()
    countries=[r["country"] for r in c.execute("SELECT country FROM zones ORDER BY country")]
    c.close()
    return jsonify({"countries":countries,"settings":get_settings()})

@app.route("/api/calculate",methods=["POST"])
def api_calculate():
    try: return jsonify(calc(request.get_json(force=True)))
    except Exception as e: return jsonify({"error":str(e)}),400

@app.route("/login",methods=["GET","POST"])
def login():
    if request.method=="POST":
        if request.form.get("password")==ADMIN_PASSWORD:
            session["admin"]=True
            return redirect(url_for("admin"))
        flash("Invalid admin password.")
    return render_template("login.html")

@app.route("/logout")
def logout():
    session.pop("admin",None)
    return redirect(url_for("index"))

@app.route("/admin")
def admin():
    if not session.get("admin"): return redirect(url_for("login"))
    c=db()
    logs=c.execute("SELECT * FROM import_log ORDER BY id DESC LIMIT 10").fetchall()
    c.close()
    return render_template("admin.html",settings=get_settings(),logs=logs)

@app.route("/admin/upload",methods=["POST"])
def upload():
    if not session.get("admin"): return redirect(url_for("login"))
    f=request.files.get("rate_file")
    if not f or not f.filename.lower().endswith(".xlsx"):
        flash("Please select an .xlsx rate master file.")
        return redirect(url_for("admin"))
    with tempfile.NamedTemporaryFile(suffix=".xlsx",delete=False) as t:
        f.save(t.name); tmp=t.name
    try:
        import_xlsx(tmp,f.filename)
        flash("Rate master updated successfully for all users.")
    except Exception as e:
        flash("Update failed: "+str(e))
    finally:
        os.unlink(tmp)
    return redirect(url_for("admin"))

@app.route("/admin/settings",methods=["POST"])
def settings_update():
    if not session.get("admin"): return redirect(url_for("login"))
    c=db()
    for k in ["exchange_rate","vat_rate","oda_fuel","oda_base_usd","oda_perkg_usd"]:
        if k in request.form:
            c.execute("INSERT INTO settings(key,value) VALUES(?,?) ON CONFLICT(key) DO UPDATE SET value=excluded.value",
                      (k,request.form[k]))
    c.commit(); c.close()
    flash("Settings updated.")
    return redirect(url_for("admin"))

if __name__=="__main__":
    init_db()
    c=db(); count=c.execute("SELECT COUNT(*) n FROM zones").fetchone()["n"]; c.close()
    if count==0 and os.path.exists(DEFAULT_XLSX):
        import_xlsx(DEFAULT_XLSX)
    app.run(host="0.0.0.0",port=int(os.environ.get("PORT",5000)),debug=False)
