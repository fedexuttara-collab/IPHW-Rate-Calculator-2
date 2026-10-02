import sqlite3, re, os
from openpyxl import load_workbook
DB='rate_master.db'; XLSX='IPHW_RATE_MASTER.xlsx'
c=sqlite3.connect(DB)
c.executescript('''CREATE TABLE IF NOT EXISTS settings(key TEXT PRIMARY KEY,value TEXT NOT NULL);CREATE TABLE IF NOT EXISTS zones(country TEXT PRIMARY KEY,zone TEXT);CREATE TABLE IF NOT EXISTS rates(calculator TEXT,weight_from REAL,weight_to REAL,zone TEXT,rate REAL,PRIMARY KEY(calculator,weight_from,weight_to,zone));CREATE TABLE IF NOT EXISTS oda(id INTEGER PRIMARY KEY AUTOINCREMENT,country TEXT,city TEXT,postal_from INTEGER,postal_to INTEGER,oda_type TEXT);CREATE TABLE IF NOT EXISTS import_log(id INTEGER PRIMARY KEY AUTOINCREMENT,imported_at TEXT DEFAULT CURRENT_TIMESTAMP,filename TEXT,status TEXT,message TEXT);''')
wb=load_workbook(XLSX,data_only=True)
ws=wb['Calculatore -A']
for k,v in {'exchange_rate':ws['B7'].value,'vat_rate':ws['B8'].value,'oda_fuel':ws['B9'].value,'oda_base_usd':ws['B15'].value,'oda_perkg_usd':ws['B16'].value}.items(): c.execute('INSERT OR REPLACE INTO settings VALUES(?,?)',(k,str(v)))
c.execute('DELETE FROM zones')
for r in range(5,wb['Sheet2'].max_row+1):
 a,b=wb['Sheet2'].cell(r,1).value,wb['Sheet2'].cell(r,2).value
 if a and b:c.execute('INSERT OR REPLACE INTO zones VALUES(?,?)',(str(a).strip(),str(b).strip()))
c.execute('DELETE FROM rates')
for sh,calc in [('Rate ; A','A'),('Rate; B','B')]:
 ws=wb[sh]; zs=[ws.cell(1,col).value for col in range(2,ws.max_column+1)]
 for r in range(2,9):
  label=ws.cell(r,1).value
  if not label:continue
  m=re.match(r'(\d+)-(\d+)',str(label)); lo,hi=(101,999999) if str(label).startswith('101') else (int(m.group(1)),int(m.group(2)))
  for i,z in enumerate(zs,2):
   v=ws.cell(r,i).value
   if z and isinstance(v,(int,float)):c.execute('INSERT INTO rates VALUES(?,?,?,?,?)',(calc,lo,hi,str(z),float(v)))
c.execute('DELETE FROM oda'); ws=wb['Sheet1']
for r in range(9,ws.max_row+1):
 country=ws.cell(r,1).value; city=ws.cell(r,2).value; pf=ws.cell(r,3).value; pt=ws.cell(r,4).value
 if not country:continue
 try:pf=int(float(pf)) if pf is not None else None
 except:pf=None
 try:pt=int(float(pt)) if pt is not None else None
 except:pt=None
 city=str(city).strip() if city is not None else ''
 typ='Postal' if pf is not None else 'City'
 c.execute('INSERT INTO oda(country,city,postal_from,postal_to,oda_type) VALUES(?,?,?,?,?)',(str(country).strip(),city,pf,pt,typ))
c.execute('INSERT INTO import_log(filename,status,message) VALUES(?,?,?)',(XLSX,'OK','Initial master import'));c.commit();c.close()
